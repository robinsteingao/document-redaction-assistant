from __future__ import annotations

import builtins
import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from redaction_assistant.external_release import write_external_release
from redaction_assistant.workflow import build_redaction_package


@pytest.mark.parametrize(
    ("suffix", "missing_module"),
    [(".docx", "lxml"), (".docx", "docx"), (".xlsx", "lxml"), (".xlsx", "openpyxl")],
)
def test_missing_office_validator_blocks_external_copy(tmp_path: Path, monkeypatch, suffix: str, missing_module: str):
    source = tmp_path / ("synthetic" + suffix)
    if suffix == ".docx":
        from docx import Document

        doc = Document()
        doc.add_paragraph("联系人：张三。技术指标：10kV。")
        doc.save(source)
    else:
        from openpyxl import Workbook

        book = Workbook()
        book.active["A1"] = "联系人：张三。技术指标：10kV。"
        book.save(source)
        book.close()
    package, mapping = build_redaction_package([source], project_alias_id="SYNTHETIC-DEPENDENCY")
    original_import = builtins.__import__

    def without_validator(name, *args, **kwargs):
        if name == missing_module or name.startswith(missing_module + "."):
            raise ImportError("synthetic missing optional validator")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", without_validator)
    paths = write_external_release(tmp_path / "out", [source], package, mapping)
    manifest = json.loads(paths["external_release_manifest"].read_text(encoding="utf-8"))
    assert manifest["exported_count"] == 0
    assert manifest["blocked_count"] == 1
    assert list(paths["external_release"].iterdir()) == [paths["external_release_manifest"]]
    assert source.is_file()
    assert package["redacted_text_blocks"]
    assert "10kV" in package["redacted_text_blocks"][0]["text"]
