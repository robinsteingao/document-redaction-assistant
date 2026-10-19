from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from redaction_assistant import hidden_channels
from redaction_assistant.external_release import write_external_release
from redaction_assistant.workflow import build_redaction_package


@pytest.mark.parametrize("evidence", ["missing", None, "", 0, [], "invalid-hash"])
def test_missing_or_invalid_mapping_binding_cannot_export_unmapped_source(tmp_path: Path, evidence):
    source = tmp_path / "source.txt"
    source.write_text("SYNTHETIC_PRIVATE_TOKEN", encoding="utf-8")
    package, _mapping = build_redaction_package(
        [source], project_alias_id="SYNTHETIC-BINDING",
        customer_dictionary={"organization": ["SYNTHETIC_PRIVATE_TOKEN"]},
    )
    if evidence == "missing":
        del package["field_mappings_hash"]
    else:
        package["field_mappings_hash"] = evidence
    outputs = write_external_release(tmp_path / "out", [source], package, {"items": []})
    manifest = json.loads(outputs["external_release_manifest"].read_text(encoding="utf-8"))
    assert manifest["exported_count"] == 0
    assert manifest["blocked_count"] == 1
    assert not list(outputs["external_release"].glob("*.txt"))


@pytest.mark.parametrize("suffix", [".docx", ".xlsx"])
def test_channel_extraction_failure_is_not_empty_clearance(tmp_path: Path, monkeypatch, suffix: str):
    source = tmp_path / ("source" + suffix)
    if suffix == ".docx":
        from docx import Document

        doc = Document()
        doc.add_paragraph("公开正文")
        doc.sections[0].header.paragraphs[0].text = "联系人：张三"
        doc.save(source)
        extractor = "_docx_texts"
    else:
        from openpyxl import Workbook
        from openpyxl.workbook.defined_name import DefinedName

        book = Workbook()
        book.active["A1"] = "公开正文"
        book.defined_names.add(DefinedName("SyntheticName", attr_text='"联系人：张三"'))
        book.save(source)
        book.close()
        extractor = "_xlsx_channel_texts"
    original = source.read_bytes()
    package, mapping = build_redaction_package([source], project_alias_id="SYNTHETIC-CHANNEL-FAILURE")

    def failed_extraction(_path):
        raise OSError("synthetic channel failure")

    monkeypatch.setattr(hidden_channels, extractor, failed_extraction)
    outputs = write_external_release(tmp_path / "out", [source], package, mapping)
    manifest = json.loads(outputs["external_release_manifest"].read_text(encoding="utf-8"))
    assert manifest["exported_count"] == 0
    assert manifest["blocked_count"] == 1
    assert not list(outputs["external_release"].glob("*" + suffix))
    assert source.read_bytes() == original
