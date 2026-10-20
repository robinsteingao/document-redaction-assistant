from unittest.mock import patch

import pytest
from docx import Document

from redaction_assistant import core_dependencies as deps
from redaction_assistant.review import build_review_workspace
from redaction_assistant.workflow import build_redaction_package


@pytest.mark.parametrize("text,kind", [("技术指标：10kV。", "technical_metric"), ("现场验证已完成", "validation_evidence"), ("合同金额：100万元", "amount")])
def test_hidden_assessment_fields_share_review_and_impact(tmp_path, text, kind):
    doc = Document()
    doc.add_paragraph("项目说明")
    doc.sections[0].header.paragraphs[0].text = text
    source = tmp_path / "input.docx"
    doc.save(source)
    review = build_review_workspace(iter([source]), project_alias_id="SYNTHETIC")
    candidate = next(c for c in review["review_candidates"] if c["kind"] == kind)
    cid = candidate["candidate_id"]
    default_package, default_mapping = build_redaction_package([source], project_alias_id="SYNTHETIC")
    default_items = [i for i in default_mapping["items"] if i["candidate_id"] == cid]
    if kind == "amount":
        assert default_items[0]["strategy"] == "range"
    else:
        assert not default_items
    for action, strategy, expected in [("keep", "keep", "pass"), ("redact", "mask", "blocked_requires_confirmation")]:
        package, mapping = build_redaction_package([source], project_alias_id="SYNTHETIC", review_decisions={cid: {"action": action, "strategy": strategy}})
        impact = package["redaction_impact_summary"]
        assert impact["overall_level"] == expected
        assert cid not in impact["unmatched_decisions"]
        assert sum(i["candidate_id"] == cid for i in impact["items"]) == 1
    package, _ = build_redaction_package([source], project_alias_id="SYNTHETIC", review_decisions={cid: {"action": "redact", "strategy": "mask"}}, customer_confirmed_degradation_risk=True)
    assert package["redaction_impact_summary"]["customer_confirmed"] is True
    doc.add_paragraph(text)
    doc.save(source)
    package, _ = build_redaction_package([source], project_alias_id="SYNTHETIC", review_decisions={cid: {"action": "redact", "strategy": "mask"}})
    impact = package["redaction_impact_summary"]
    assert impact["overall_level"] == "blocked_requires_confirmation"
    assert sum(i["candidate_id"] == cid for i in impact["items"]) == 1


def test_dependency_versions_match_declared_bounds():
    from pathlib import Path
    requirements = (Path(__file__).resolve().parents[1] / "requirements.txt").read_text(encoding="utf-8")
    assert "cryptography>=41.0.0,<43" in requirements
    assert deps.REQUIREMENTS["cryptography"] == ((41, 0, 0), (43, 0, 0))


@pytest.mark.parametrize("version,expected", [("40.0.0", 2), ("41.0.0", 0), ("42.0.8", 0), ("43.0.0", 2), ("46.0.4", 2), ("42.0.0rc1", 2), (None, 1)])
def test_core_dependency_check(version, expected):
    def installed(name):
        if name == "cryptography" and version is None:
            raise deps.importlib.metadata.PackageNotFoundError(name)
        return {"cryptography": version, "pypdf": "4.0.0", "defusedxml": "0.7.1"}[name]
    with patch.object(deps.importlib.metadata, "version", side_effect=installed), patch.object(deps.importlib, "import_module"):
        assert deps.check_core_dependencies()[0] == expected


def test_old_python_is_rejected():
    with patch.object(deps.sys, "version_info", (3, 9)), patch.object(deps.importlib.metadata, "version", return_value="42.0.8"), patch.object(deps.importlib, "import_module"):
        assert deps.check_core_dependencies()[0] == 2
