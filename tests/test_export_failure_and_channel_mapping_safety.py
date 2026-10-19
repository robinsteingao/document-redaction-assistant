from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from redaction_assistant.crypto import decrypt_mapping_file
from redaction_assistant.external_release import write_external_release
from redaction_assistant.local_service import handle_request
from redaction_assistant.redactor import mapping_hash, restore_text
from redaction_assistant.workflow import build_redaction_package, write_package


def _plain_source(tmp_path: Path) -> Path:
    source = tmp_path / "input.txt"
    source.write_text("联系人：王五。技术指标：10kV。", encoding="utf-8")
    return source


def _docx_with_header(tmp_path: Path, body: str, header: str) -> Path:
    document = Document()
    document.add_paragraph(body)
    document.sections[0].header.paragraphs[0].text = header
    path = tmp_path / "report.docx"
    document.save(path)
    return path


def _occupied_output(tmp_path: Path) -> Path:
    out = tmp_path / "out"
    out.mkdir()
    # 同名普通文件占位：令外发批次目录创建失败，稳定复现导出失败路径。
    (out / "external_release").write_text("occupied", encoding="utf-8")
    return out


def _read_all_text(root: Path) -> str:
    return "\n".join(
        path.read_text(encoding="utf-8", errors="ignore")
        for path in sorted(root.rglob("*"))
        if path.is_file()
    )


def _exported_item(outputs: dict) -> dict:
    manifest = json.loads(outputs["external_release_manifest"].read_text(encoding="utf-8"))
    return next(item for item in manifest["files"] if item["status"] == "exported")


def test_write_package_failure_leaves_no_mapping_and_no_plaintext(tmp_path: Path):
    source = _plain_source(tmp_path)
    out = _occupied_output(tmp_path)
    package, mapping = build_redaction_package([source], project_alias_id="FAIL-CLOSED")

    with pytest.raises(Exception):
        write_package(out, package, mapping, source_files=[source], mapping_passphrase="pw-synthetic-123")

    assert not list(out.rglob("local_mapping*"))
    assert "王五" not in _read_all_text(out)


def test_service_build_failure_leaves_no_mapping_and_no_plaintext(tmp_path: Path):
    source = _plain_source(tmp_path)
    out = _occupied_output(tmp_path)

    response = handle_request({
        "action": "build_package",
        "project_alias_id": "FAIL-CLOSED-SERVICE",
        "out": str(out),
        "files": [str(source)],
        "mapping_passphrase": "pw-synthetic-123",
    })

    assert response["success"] is False
    assert not list(out.rglob("local_mapping*"))
    assert "王五" not in _read_all_text(out)


def test_body_and_header_persons_share_unique_persistent_mapping(tmp_path: Path):
    source = _docx_with_header(tmp_path, "联系人：王五，负责试运行。", "联系人：张三")
    out = tmp_path / "out"
    package, mapping = build_redaction_package([source], project_alias_id="CHANNEL-UNION")

    persons = {item["original"]: item["placeholder"] for item in mapping["items"] if item["kind"] == "person"}
    assert {"王五", "张三"} <= set(persons)
    assert persons["王五"] != persons["张三"]
    assert len(set(persons.values())) == len(persons)
    assert package["field_mappings_hash"] == mapping_hash(mapping)

    outputs = write_package(out, package, mapping, source_files=[source], mapping_passphrase="pw-synthetic-123")
    saved = decrypt_mapping_file(outputs["encrypted_mapping"], passphrase="pw-synthetic-123")
    saved_by_original = {item["original"]: item["placeholder"] for item in saved["items"]}
    assert saved_by_original["王五"] == persons["王五"]
    assert saved_by_original["张三"] == persons["张三"]

    exported = outputs["external_release"] / _exported_item(outputs)["output_name"]
    document = Document(exported)
    body_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    header_text = "\n".join(paragraph.text for paragraph in document.sections[0].header.paragraphs)
    assert "王五" not in body_text and "张三" not in header_text
    assert saved_by_original["王五"] in body_text
    assert saved_by_original["张三"] in header_text
    assert restore_text(header_text, saved).strip() == "联系人：张三"
    assert "联系人：王五" in restore_text(body_text, saved)


def test_same_person_in_body_and_header_uses_single_item(tmp_path: Path):
    source = _docx_with_header(tmp_path, "联系人：王五，负责试运行。", "联系人：王五")

    _package, mapping = build_redaction_package([source], project_alias_id="CHANNEL-DEDUP")

    persons = [item for item in mapping["items"] if item["kind"] == "person"]
    assert len(persons) == 1
    assert persons[0]["original"] == "王五"


def test_placeholders_stay_unique_across_files(tmp_path: Path):
    first = tmp_path / "first.txt"
    first.write_text("联系人：王五。", encoding="utf-8")
    second = tmp_path / "second.txt"
    second.write_text("联系人：赵六。", encoding="utf-8")

    _package, mapping = build_redaction_package([first, second], project_alias_id="MULTI-FILE")

    persons = {item["original"]: item["placeholder"] for item in mapping["items"] if item["kind"] == "person"}
    assert {"王五", "赵六"} <= set(persons)
    assert persons["王五"] != persons["赵六"]


def test_export_fails_closed_when_mapping_misses_channel_entity(tmp_path: Path):
    source = _plain_source(tmp_path)
    package, _mapping = build_redaction_package([source], project_alias_id="MAPPING-DRIFT")
    stripped = {"items": []}
    package["field_mappings_hash"] = mapping_hash(stripped)

    outputs = write_external_release(tmp_path / "out", [source], package, stripped)

    manifest = json.loads(outputs["external_release_manifest"].read_text(encoding="utf-8"))
    assert manifest["exported_count"] == 0
    assert manifest["blocked_count"] == 1
    assert manifest["files"][0]["reason"] == "channel_mapping_drift"
    assert not list(outputs["external_release"].glob("*.txt"))
