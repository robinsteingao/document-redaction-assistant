from __future__ import annotations

import json
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from redaction_assistant import external_release as release
from redaction_assistant.parsers import parse_file
from redaction_assistant.redactor import candidate_id_for, mapping_hash
from redaction_assistant.rules import detect_entities
from redaction_assistant.workflow import build_redaction_package


MAPPING = {"items": [{"original": "张三", "replacement": "人员A", "kind": "person"}]}


def package_for(source, mapping=None):
    _, item = parse_file(source)
    return {
        "residual_scan": {"status": "clear", "hit_count": 0},
        "source_file_manifest": [item],
        "field_mappings_hash": mapping_hash(mapping or MAPPING),
    }


def export(tmp_path, source, package=None, **kwargs):
    # 与产品路径一致：导出前先把通道内新实体并入映射；导出本身不再临时发明占位符。
    mapping = release.augment_mapping_for_export(
        [source], MAPPING, customer_dictionary=None, review_decisions=None,
    )
    if package is None:
        package = package_for(source, mapping)
    paths = release.write_external_release(tmp_path / "out", [source], package, mapping, **kwargs)
    manifest = json.loads(paths["external_release_manifest"].read_text(encoding="utf-8"))
    return paths, manifest["files"][0]


def change_zip(source, updates):
    with zipfile.ZipFile(source) as zf:
        payload = {name: zf.read(name) for name in zf.namelist()}
    payload.update(updates)
    with zipfile.ZipFile(source, "w") as zf:
        for name, data in payload.items():
            zf.writestr(name, data)


@pytest.mark.parametrize("status", [None, "not_scanned", "residual_found", "CLEAR", "unknown"])
def test_residual_must_be_explicit_clear(tmp_path, status):
    source = tmp_path / "source.txt"
    source.write_text("联系人：张三", encoding="utf-8")
    package = package_for(source)
    package["residual_scan"] = {} if status is None else {"status": status}
    paths, item = export(tmp_path, source, package)
    assert item["status"] == "blocked"
    assert not list(paths["external_release"].glob("*.txt"))


@pytest.mark.parametrize("mutation", ["hash", "missing", "failed", "hit_count"])
def test_source_and_parse_evidence_must_match(tmp_path, mutation):
    source = tmp_path / "source.txt"
    source.write_text("联系人：张三", encoding="utf-8")
    package = package_for(source)
    if mutation == "hash":
        source.write_text("联系人：李四", encoding="utf-8")
    elif mutation == "missing":
        del package["source_file_manifest"]
    elif mutation == "failed":
        package["source_file_manifest"][0]["parser_status"] = "failed"
    else:
        package["residual_scan"]["hit_count"] = 1
    _, item = export(tmp_path, source, package)
    assert item["status"] == "blocked"


def test_old_batches_and_private_named_user_files_are_untouched(tmp_path):
    source = tmp_path / "张三-local_mapping-registration.txt"
    source.write_text("联系人：张三", encoding="utf-8")
    old = tmp_path / "out" / "external_release" / "old" / "license.json"
    old.parent.mkdir(parents=True)
    old.write_text("user-owned", encoding="utf-8")
    first, item = export(tmp_path, source)
    second, _ = export(tmp_path, source)
    assert first["external_release"] != second["external_release"]
    assert first["external_release"].parent == old.parent.parent
    assert old.read_text() == "user-owned"
    for paths in (first, second):
        manifest = json.loads(paths["external_release_manifest"].read_text(encoding="utf-8"))
        names = {record["output_name"] for record in manifest["files"] if record["status"] == "exported"}
        assert set(p.name for p in paths["external_release"].iterdir()) == names | {"external_release_manifest.json"}
        assert "张三" not in json.dumps(manifest, ensure_ascii=False)
        assert all("local_mapping" not in name and "registration" not in name for name in names)
        assert (paths["external_release"] / item["output_name"]).read_text(encoding="utf-8") == "联系人：人员A"


@pytest.mark.parametrize("bad_output", ["missing", "empty", "corrupt", "residual"])
def test_generated_office_must_exist_and_reopen(tmp_path, monkeypatch, bad_output):
    source = tmp_path / "document.docx"
    doc = Document()
    doc.add_paragraph("联系人：张三")
    doc.save(source)

    def broken(src, dest, mapping):
        if bad_output == "empty":
            dest.write_bytes(b"")
        elif bad_output == "corrupt":
            dest.write_bytes(b"not-office")
        elif bad_output == "residual":
            dest.write_bytes(src.read_bytes())

    monkeypatch.setattr(release, "_rewrite_docx", broken)
    paths, item = export(tmp_path, source)
    assert item["status"] == "blocked"
    assert not list(paths["external_release"].glob("*.docx"))


def test_word_span_redaction_preserves_runs_and_qname_namespaces(tmp_path):
    source = tmp_path / "styled.docx"
    doc = Document()
    para = doc.add_paragraph()
    para.add_run("保留前缀 ").bold = True
    para.add_run("张").italic = True
    para.add_run("三").underline = True
    tail = para.add_run(" 技术指标10kV 后缀")
    tail.font.size = Pt(18)
    doc.save(source)
    paths, item = export(tmp_path, source)
    assert item["status"] == "exported"
    dest = paths["external_release"] / item["output_name"]
    result = Document(dest)
    runs = result.paragraphs[0].runs
    assert result.paragraphs[0].text == "保留前缀 人员A 技术指标10kV 后缀"
    assert runs[0].text == "保留前缀 " and runs[0].bold
    assert runs[1].text == "人员A" and runs[1].italic
    assert runs[2].text == "" and runs[2].underline
    assert runs[3].text == " 技术指标10kV 后缀" and runs[3].font.size == Pt(18)
    from lxml import etree
    with zipfile.ZipFile(dest) as zf:
        assert not any(name.startswith("docProps/thumbnail.") for name in zf.namelist())
        assert b"thumbnail" not in zf.read("_rels/.rels")
        root = etree.fromstring(zf.read("word/document.xml"))
        for node in root.iter():
            for attr, value in node.attrib.items():
                if attr.endswith("}Ignorable"):
                    assert all(prefix in node.nsmap for prefix in value.split())


def test_xlsx_reopens_and_preserves_styles_formulas_and_numbers(tmp_path):
    source = tmp_path / "styled.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "张三"
    ws["A1"].font = Font(bold=True, color="FF0000")
    ws["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["B1"] = 123.45
    ws["B1"].number_format = "0.00"
    ws["C1"] = "=B1*2"
    ws.column_dimensions["A"].width = 28
    wb.save(source)
    paths, item = export(tmp_path, source)
    assert item["status"] == "exported"
    result = load_workbook(paths["external_release"] / item["output_name"])
    sheet = result.active
    assert sheet["A1"].value == "人员A"
    assert sheet["A1"].font.bold and sheet["A1"].font.color.rgb == "00FF0000"
    assert sheet["A1"].fill.fgColor.rgb == "00FFFF00"
    assert sheet["B1"].value == 123.45 and sheet["B1"].number_format == "0.00"
    assert sheet["C1"].value == "=B1*2"
    assert sheet.column_dimensions["A"].width == 28
    result.close()


@pytest.mark.parametrize("part", ["word/media/image1.png", "word/embeddings/object.bin", "word/unknown.xml"])
def test_stale_hidden_clear_cannot_bypass_opaque_parts(tmp_path, part):
    source = tmp_path / "opaque.docx"
    doc = Document()
    doc.add_paragraph("公开正文")
    doc.save(source)
    change_zip(source, {part: b"<secret>private-image-or-object</secret>"})
    package = package_for(source)
    package["hidden_channel_report"] = {"files": [{"source_index": 0, "status": "clear", "channels": []}]}
    _, item = export(tmp_path, source, package)
    assert item["status"] == "blocked"


def test_xml_attributes_and_revision_authors_are_cleaned(tmp_path):
    source = tmp_path / "attribute.docx"
    doc = Document()
    doc.add_paragraph("张三")
    doc.save(source)
    with zipfile.ZipFile(source) as zf:
        xml = zf.read("word/document.xml").decode()
    xml = xml.replace("<w:p>", '<w:p><w:ins w:id="1" w:author="任意未识别作者" w:date="2026-01-01T00:00:00Z"><w:r><w:t>公开修订</w:t></w:r></w:ins><w:bookmarkStart w:id="9" w:name="张三"/>', 1)
    change_zip(source, {"word/document.xml": xml})
    paths, item = export(tmp_path, source)
    assert item["status"] == "exported"
    with zipfile.ZipFile(paths["external_release"] / item["output_name"]) as zf:
        data = zf.read("word/document.xml").decode()
        assert "张三" not in data and "任意未识别作者" not in data


def test_external_relationship_is_blocked_even_with_stale_report(tmp_path):
    source = tmp_path / "link.docx"
    doc = Document()
    doc.add_paragraph("公开正文")
    doc.save(source)
    rels = '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rSecret" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.invalid/张三" TargetMode="External"/></Relationships>'
    change_zip(source, {"word/_rels/document.xml.rels": rels})
    _, item = export(tmp_path, source)
    assert item["status"] == "blocked"
    assert item["reason"] == "external_relationship"


@pytest.mark.parametrize("kind", ["text", "mixed_scan", "rotated", "malformed"])
def test_every_pdf_is_blocked_without_invoking_engine(tmp_path, kind):
    from io import BytesIO
    from PIL import Image, ImageDraw
    from pypdf import PdfReader, PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    source = tmp_path / (kind + ".pdf")
    writer = PdfWriter()
    page = writer.add_blank_page(100, 100)
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 10 Tf 10 50 Td (SyntheticPrivate) Tj ET")
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    page[NameObject("/Contents")] = writer._add_object(stream)
    if kind == "rotated":
        page.rotate(90)
    if kind == "mixed_scan":
        image = Image.new("RGB", (100, 100), "white")
        ImageDraw.Draw(image).text((5, 50), "SCAN_PRIVATE", fill="black")
        scanned = BytesIO()
        image.save(scanned, format="PDF")
        writer.add_page(PdfReader(BytesIO(scanned.getvalue())).pages[0])
    writer.write(source)
    if kind == "malformed":
        source.write_bytes(b"%PDF-broken")

    def forbidden(*args):
        pytest.fail("PDF engine must never be called")

    paths, item = export(tmp_path, source, pdf_engine=forbidden)
    assert item["status"] == "blocked" and item["reason"] == "pdf_external_release_disabled"
    assert not list(paths["external_release"].glob("*.pdf"))


@pytest.mark.parametrize("channel", ["attribute", "relationship", "split_runs"])
def test_final_residual_check_includes_attributes_rels_and_split_runs(tmp_path, monkeypatch, channel):
    from lxml import etree

    source = tmp_path / "residual.docx"
    doc = Document()
    doc.add_paragraph("联系人：张三")
    doc.save(source)
    original_rewrite = release._rewrite_docx

    def reintroduce(src, dest, mapping):
        original_rewrite(src, dest, mapping)
        part = "_rels/.rels" if channel == "relationship" else "word/document.xml"
        with zipfile.ZipFile(dest) as zf:
            root = etree.fromstring(zf.read(part))
        if channel == "attribute":
            root.set("private", "张三")
        elif channel == "relationship":
            root[0].set("private", "张三")
        else:
            paragraph = next(root.iter(release.W_P))
            for child in list(paragraph):
                paragraph.remove(child)
            for value in ("联系人：张", "三"):
                run = etree.SubElement(paragraph, f"{{{release.W_NS}}}r")
                etree.SubElement(run, release.W_T).text = value
        change_zip(dest, {part: etree.tostring(root)})

    monkeypatch.setattr(release, "_rewrite_docx", reintroduce)
    paths, item = export(tmp_path, source)
    assert item["status"] == "blocked" and item["reason"] == "residual_after_rewrite"
    assert not list(paths["external_release"].glob("*.docx"))


def test_amount_technical_metrics_and_explicit_keep_decisions_are_unchanged(tmp_path):
    source = tmp_path / "keep.docx"
    text = "联系人：张三。投入150万元。技术指标：10kV试运行30天。"
    doc = Document()
    doc.add_paragraph(text)
    doc.save(source)
    person = next(entity for entity in detect_entities([text]) if entity.kind == "person")
    decisions = {candidate_id_for(person): {"action": "keep"}}
    package, mapping = build_redaction_package([source], project_alias_id="SYNTHETIC-KEEP", review_decisions=decisions)
    paths = release.write_external_release(tmp_path / "out", [source], package, mapping, review_decisions=decisions)
    item = json.loads(paths["external_release_manifest"].read_text(encoding="utf-8"))["files"][0]
    assert item["status"] == "exported", item
    output = Document(paths["external_release"] / item["output_name"]).paragraphs[0].text
    assert output == package["redacted_text_blocks"][0]["text"]
    assert "张三" in output and "10kV试运行30天" in output
    assert "150万元" not in output and "100万至500万" in output


def test_failed_atomic_publish_does_not_touch_old_batch(tmp_path, monkeypatch):
    source = tmp_path / "source.txt"
    source.write_text("公开", encoding="utf-8")
    old, _ = export(tmp_path, source)
    contents = {p.name: p.read_bytes() for p in old["external_release"].iterdir()}
    original = Path.rename

    def deny(self, target):
        if Path(target).parent.name == "external_release":
            raise OSError("simulated publish failure")
        return original(self, target)

    monkeypatch.setattr(Path, "rename", deny)
    with pytest.raises(OSError):
        export(tmp_path, source)
    assert {p.name: p.read_bytes() for p in old["external_release"].iterdir()} == contents
    assert list(old["external_release"].parent.iterdir()) == [old["external_release"]]
