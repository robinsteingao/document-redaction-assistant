from __future__ import annotations

import hashlib
import json
import posixpath
import re
import tempfile
import uuid
import zipfile
from pathlib import Path
from typing import Any, Callable, Iterable
from urllib.parse import unquote, urlsplit

from .hidden_channels import extract_channel_texts, inspect_hidden_channels, ooxml_safety_channels
from .parsers import _check_zip_budget, file_sha256, parse_file
from .redactor import _collect_spans, build_mapping, mapping_hash, redact_text
from .rules import detect_entities_with_dictionary


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
X_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
W_P = f"{{{W_NS}}}p"
W_T = f"{{{W_NS}}}t"
X_T = f"{{{X_NS}}}t"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
BLOCKING_KINDS = {"embedded_object", "binary_media", "unknown_part", "external_relationship",
                  "hidden_sheet", "very_hidden_sheet", "comment"}
PdfEngine = Callable[[Path, Path, list[str]], dict[str, str]]


class ExportBlocked(ValueError):
    """Only fixed, non-sensitive reason codes may cross the export boundary."""


def write_external_release(
    output_dir: Path | str,
    source_files: Iterable[Path | str],
    package: dict[str, Any],
    mapping: dict[str, Any],
    *,
    customer_dictionary: Path | str | dict | None = None,
    review_decisions: dict[str, dict[str, Any]] | None = None,
    pdf_engine: PdfEngine | None = None,
) -> dict[str, Path]:
    """Publish only a validated new batch. Old batches and user files are immutable.

    The temporary workspace is owned by this invocation. Source snapshots stay
    outside its publish subtree and are never included in the external manifest.
    PDF export is intentionally disabled, including custom engines.
    """
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    parent = output / "external_release"
    if parent.is_symlink():
        raise ExportBlocked("unsafe_release_directory")
    parent.mkdir(exist_ok=True)
    job = uuid.uuid4().hex
    release_dir = parent / job
    files = [Path(path) for path in source_files]
    with tempfile.TemporaryDirectory(prefix=f".export-{job}-", dir=output) as workspace:
        workspace = Path(workspace)
        inputs = workspace / "inputs"
        inputs.mkdir()
        staged = workspace / "publish"
        staged.mkdir()
        results = []
        for index, source in enumerate(files):
            suffix = source.suffix.lower()
            alias = f"document_{index + 1:04d}" + (suffix if suffix in {".txt", ".md", ".docx", ".xlsx", ".pdf"} else "")
            if suffix == ".pdf":
                results.append(_blocked(alias, "pdf_external_release_disabled"))
                continue
            try:
                _check_package_gate(package, mapping, len(files))
                data = source.read_bytes()
                digest = hashlib.sha256(data).hexdigest()
                evidence = package["source_file_manifest"][index]
                if (not isinstance(evidence, dict) or evidence.get("file_name") != source.name
                        or evidence.get("file_hash") != digest):
                    raise ExportBlocked("source_hash_mismatch")
                if evidence.get("parser_status") != "ok":
                    raise ExportBlocked("source_parse_not_ok")
                if any("ocr_low_confidence" in str(w) for w in evidence.get("warnings", [])):
                    raise ExportBlocked("source_parse_unverified")
                snapshot = inputs / alias
                snapshot.write_bytes(data)
                _, fresh_parse = parse_file(snapshot)
                if fresh_parse.get("parser_status") != "ok" or fresh_parse.get("file_hash") != digest:
                    raise ExportBlocked("source_parse_not_ok")
                # Never trust a cached hidden report, even if its status says clear.
                hidden = inspect_hidden_channels([snapshot])["files"][0]
                export_mapping = _mapping_with_channel_text(
                    [snapshot], mapping, customer_dictionary=customer_dictionary,
                    review_decisions=review_decisions,
                )
                result = _export_one(snapshot, staged, export_mapping, _needles(export_mapping), hidden)
                if result["status"] == "exported":
                    result["source_sha256"] = digest
                results.append(result)
            except ExportBlocked as exc:
                results.append(_blocked(alias, str(exc)))
            except Exception:
                results.append(_blocked(alias, "source_unverified"))
        manifest = {
            "schema_version": "external_release.v1",
            "batch_id": job,
            "direct_release_claim": False,
            "note": "仅校验已识别且应替换字段，不声称识别全部敏感内容；仍需人工复核。PDF外发暂时禁用。",
            "files": results,
            "exported_count": sum(item["status"] == "exported" for item in results),
            "blocked_count": sum(item["status"] == "blocked" for item in results),
        }
        manifest_name = "external_release_manifest.json"
        (staged / manifest_name).write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        # Exact allowlist: no source snapshots, private mapping or unlisted remnants.
        expected = {item["output_name"] for item in results if item["status"] == "exported"}
        if {p.name for p in staged.iterdir()} != expected | {manifest_name}:
            raise ExportBlocked("unexpected_batch_file")
        for item in results:
            if item["status"] == "exported":
                path = staged / item["output_name"]
                _validate_output(path)
                if file_sha256(path) != item["output_sha256"]:
                    raise ExportBlocked("output_changed_after_validation")
        json.loads((staged / manifest_name).read_text(encoding="utf-8"))
        # Same filesystem atomic directory rename; never replace an existing job.
        if release_dir.exists():
            raise ExportBlocked("batch_already_exists")
        staged.rename(release_dir)
    return {"external_release": release_dir, "external_release_manifest": release_dir / manifest_name}


def _check_package_gate(package: dict[str, Any], mapping: dict[str, Any], count: int) -> None:
    scan = package.get("residual_scan")
    if (not isinstance(scan, dict) or scan.get("status") != "clear"
            or type(scan.get("hit_count")) is not int or scan["hit_count"] != 0):
        raise ExportBlocked("residual_not_clear")
    evidence = package.get("source_file_manifest")
    if not isinstance(evidence, list) or len(evidence) != count:
        raise ExportBlocked("source_evidence_missing")
    expected_mapping = package.get("field_mappings_hash")
    if (not isinstance(expected_mapping, str)
            or re.fullmatch(r"[0-9a-f]{64}", expected_mapping) is None
            or expected_mapping != mapping_hash(mapping)):
        raise ExportBlocked("mapping_hash_mismatch")


def _export_one(source, release_dir, mapping, needles, file_report) -> dict[str, Any]:
    dest = release_dir / source.name  # source is a private snapshot with a safe alias
    try:
        if file_report.get("status") not in {"clear", "found"}:
            raise ExportBlocked("hidden_channels_unverified")
        kinds = {channel.get("kind") for channel in file_report.get("channels", [])}
        blocked = kinds & (BLOCKING_KINDS - ({"comment"} if source.suffix == ".docx" else set()))
        if blocked:
            raise ExportBlocked("external_relationship" if "external_relationship" in blocked else sorted(blocked)[0])
        if source.suffix == ".docx":
            _rewrite_docx(source, dest, mapping)
        elif source.suffix == ".xlsx":
            _rewrite_xlsx(source, dest, mapping)
        elif source.suffix in {".txt", ".md"}:
            dest.write_text(redact_text(source.read_text(encoding="utf-8"), mapping), encoding="utf-8")
        else:
            raise ExportBlocked("unsupported")
        _validate_output(dest)
        if _contains_needle(dest, needles):
            raise ExportBlocked("residual_after_rewrite")
        return {"file_name": source.name, "status": "exported", "output_name": dest.name,
                "reason": "", "output_sha256": file_sha256(dest)}
    except ExportBlocked as exc:
        _discard(dest)
        return _blocked(source.name, str(exc))
    except Exception:
        _discard(dest)
        return _blocked(source.name, "rewrite_or_validation_failed")


def _xml_root(data: bytes):
    # lxml preserves in-scope prefixes used only in QName attribute VALUES
    # (notably mc:Ignorable), unlike ElementTree's namespace reserialization.
    from lxml import etree
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False,
                             remove_comments=True, remove_pis=True)
    root = etree.fromstring(data, parser)
    if root.getroottree().docinfo.doctype:
        raise ExportBlocked("xml_dtd_unsupported")
    for node in root.iter():
        if not isinstance(node.tag, str):
            raise ExportBlocked("xml_entity_unsupported")
        for attr, value in node.attrib.items():
            if attr == f"{{{MC_NS}}}Ignorable":
                if any(prefix not in node.nsmap for prefix in value.split()):
                    raise ExportBlocked("undeclared_qname_prefix")
    return root


def _read_ooxml(source):
    with zipfile.ZipFile(source) as zf:
        channels = ooxml_safety_channels(zf)
        blocked = {c["kind"] for c in channels} & BLOCKING_KINDS
        if blocked:
            raise ExportBlocked("external_relationship" if "external_relationship" in blocked else sorted(blocked)[0])
        payload = {name: zf.read(name) for name in zf.namelist() if not name.endswith("/")}
    thumbnails = {name for name in payload if re.fullmatch(r"docProps/thumbnail\.(jpeg|jpg|png|wmf|emf)", name)}
    for name in thumbnails:
        del payload[name]
    roots = {name: _xml_root(data) for name, data in payload.items()}
    for name, root in roots.items():
        if name.endswith(".rels"):
            for relation in list(root):
                target = _relationship_target(name, relation.attrib.get("Target", ""))
                if target in thumbnails:
                    root.remove(relation)
                elif target not in payload:
                    raise ExportBlocked("relationship_target_missing")
        elif name == "[Content_Types].xml":
            for node in list(root):
                if node.attrib.get("PartName", "").lstrip("/") in thumbnails:
                    root.remove(node)
                elif (node.tag.endswith("}Default") and node.attrib.get("Extension", "").lower()
                      in {"jpeg", "jpg", "png", "wmf", "emf"}):
                    root.remove(node)
    return roots


def _relationship_target(rels_name, target):
    parsed = urlsplit(target)
    if parsed.scheme or parsed.netloc or parsed.query or parsed.fragment or "\\" in target:
        raise ExportBlocked("external_relationship")
    if rels_name == "_rels/.rels":
        base = ""
    elif "/_rels/" in rels_name:
        base = rels_name.rsplit("/_rels/", 1)[0]
    else:
        raise ExportBlocked("unknown_relationship_part")
    resolved = posixpath.normpath(posixpath.join(base, unquote(parsed.path))) if not target.startswith("/") else unquote(parsed.path).lstrip("/")
    if resolved.startswith("../") or resolved in {".", ".."}:
        raise ExportBlocked("unsafe_relationship_target")
    return resolved


def _rewrite_docx(source, dest, mapping):
    roots = _read_ooxml(source)
    for name, root in roots.items():
        if name.startswith("docProps/") or name.startswith("word/comments"):
            _wipe_text(root)
        if name.startswith("word/") and not name.endswith(".rels"):
            _redact_word_story(root, mapping)
        else:
            _redact_text_nodes(root, mapping)
    _write_roots(dest, roots)


def _rewrite_xlsx(source, dest, mapping):
    roots = _read_ooxml(source)
    for name, root in roots.items():
        if name.startswith("docProps/"):
            _wipe_text(root)
        for sheet in root.iter(f"{{{X_NS}}}sheet"):
            if sheet.attrib.get("state") in {"hidden", "veryHidden"}:
                raise ExportBlocked("hidden_sheet")
        for si in root.iter(f"{{{X_NS}}}si"):
            _redact_joined(list(si.iter(X_T)), mapping)
        for cell in root.iter(f"{{{X_NS}}}c"):
            if cell.attrib.get("t") == "inlineStr":
                _redact_joined(list(cell.iter(X_T)), mapping)
        _redact_text_nodes(root, mapping, skip_text_tags={X_T})
    _write_roots(dest, roots)


def _redact_word_story(root, mapping):
    for node in root.iter():
        local = node.tag.rsplit("}", 1)[-1]
        if local in {"del", "moveFrom"}:
            _wipe_text(node)
        if local in {"comment", "ins", "del", "moveFrom", "moveTo"}:
            for attr in list(node.attrib):
                if attr.rsplit("}", 1)[-1] in {"author", "initials", "date", "dateUtc"}:
                    node.attrib[attr] = ""
        if local in {"instrText", "fldSimple", "altChunk"}:
            raise ExportBlocked("active_word_field_unsupported")
    for paragraph in root.iter(W_P):
        nodes = []
        for node in paragraph.iter(W_T):
            # A text box paragraph must not be counted again in its outer paragraph.
            if next((p for p in node.iterancestors() if p.tag == W_P), None) is paragraph:
                nodes.append(node)
        _redact_joined(nodes, mapping)
    _redact_text_nodes(root, mapping, skip_text_tags={W_T})


def _redact_joined(nodes, mapping):
    """Apply original-coordinate spans only to affected text runs, never merge runs."""
    texts = [node.text or "" for node in nodes]
    pairs = [(item.get("original") or "", item.get("replacement") or item.get("placeholder") or "")
             for item in sorted(mapping.get("items", []), key=lambda i: -len(i.get("original", "")))]
    spans = _collect_spans("".join(texts), pairs)
    offsets = []
    offset = 0
    for text in texts:
        offsets.append(offset)
        offset += len(text)
    for start, end, replacement in sorted(spans, reverse=True):
        first = True
        for index, original in enumerate(texts):
            left, right = offsets[index], offsets[index] + len(original)
            if start < right and end > left:
                text = nodes[index].text or ""
                nodes[index].text = text[:max(0, start - left)] + (replacement if first else "") + text[min(len(original), end - left):]
                first = False
                nodes[index].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _redact_text_nodes(root, mapping, *, skip_text_tags=frozenset()):
    for node in root.iter():
        if node.text and node.tag not in skip_text_tags:
            node.text = redact_text(node.text, mapping)
        if node.tail:
            node.tail = redact_text(node.tail, mapping)
        for attr, value in list(node.attrib.items()):
            node.attrib[attr] = redact_text(value, mapping)


def _wipe_text(root):
    for node in root.iter():
        if node.text:
            node.text = ""
        if node.tail and node.tail.strip():
            node.tail = ""
        for attr in list(node.attrib):
            if attr.rsplit("}", 1)[-1] in {"author", "initials", "date", "dateUtc", "name"}:
                node.attrib[attr] = ""


def _write_roots(dest, roots):
    from lxml import etree
    with zipfile.ZipFile(dest, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for name, root in roots.items():
            zf.writestr(name, etree.tostring(root, encoding="utf-8", xml_declaration=True))


def _validate_output(path):
    if not path.is_file() or path.is_symlink() or path.stat().st_size == 0:
        raise ExportBlocked("output_missing_or_empty")
    if path.suffix == ".docx":
        from docx import Document
        Document(path)
        _read_ooxml(path)
    elif path.suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, data_only=False)
        try:
            for sheet in workbook:
                for row in sheet:
                    for cell in row:
                        _ = cell.value
        finally:
            workbook.close()
        _read_ooxml(path)
    elif path.suffix in {".txt", ".md"}:
        path.read_text(encoding="utf-8")
    else:
        raise ExportBlocked("output_type_unsupported")


def _contains_needle(path, needles):
    # Failure is not clearance: this is always inside a fail-closed boundary.
    if not path.is_file() or path.stat().st_size == 0:
        raise ExportBlocked("output_missing_or_empty")
    if path.suffix in {".docx", ".xlsx"}:
        text = _zip_xml_text(path)
    elif path.suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8")
    else:
        raise ExportBlocked("output_type_unsupported")
    return any(needle and needle in text for needle in needles)


def _zip_xml_text(path):
    chunks = []
    with zipfile.ZipFile(path) as zf:
        _check_zip_budget(zf)
        for name in zf.namelist():
            chunks.append(name)
            root = _xml_root(zf.read(name))
            chunks.append("".join(root.itertext()))  # includes cross-run values
            for node in root.iter():
                chunks.extend(node.attrib.values())
                chunks.extend(node.nsmap.values())
    return "\n".join(chunks)


def _mapping_with_channel_text(files, mapping, *, customer_dictionary, review_decisions):
    texts = [text for path in files for text in extract_channel_texts(path)]
    if files[0].suffix in {".txt", ".md"}:
        texts.append(files[0].read_text(encoding="utf-8"))
    extra = build_mapping(detect_entities_with_dictionary(texts, customer_dictionary), review_decisions=review_decisions)
    items = list(mapping.get("items") or [])
    seen = {item.get("original") for item in items}
    for item in extra.get("items") or []:
        if item.get("original") and item["original"] not in seen:
            items.append(item)
            seen.add(item["original"])
    return {**mapping, "items": items}


def _needles(mapping):
    return sorted({str(item["original"]) for item in mapping.get("items", []) if item.get("original")}, key=len, reverse=True)


def _blocked(file_name, reason):
    return {"file_name": file_name, "status": "blocked", "output_name": "", "reason": reason}


def _discard(path):
    # Only a single generated path in this invocation's private workspace.
    if path.is_file():
        path.unlink()
